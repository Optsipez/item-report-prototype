import openpyxl, sys, pickle, datetime, time
src, out = sys.argv[1], sys.argv[2]
t = time.time()
wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
D = {}
for name in wb.sheetnames:
    rows = []
    for r in wb[name].iter_rows(values_only=True):
        rows.append(r)
    # trim trailing all-empty rows
    while rows and all(v is None or v == '' for v in rows[-1]): rows.pop()
    D[name] = rows
    print(name, len(rows), round(time.time() - t), 's', flush=True)
pickle.dump(D, open(out, 'wb'))
print('done')
